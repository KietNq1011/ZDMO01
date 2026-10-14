# Sinh template ZDMO01 (.xlsx, engine XLSX) + Sample JSON — placeholder khớp
# ZCL_ZDMO01_SOURCE=>ty_form. Usage: python build_zdmo01_xlsx.py <out_dir>
import sys, json, pathlib, zipfile, io
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

out = pathlib.Path(sys.argv[1])

wb = Workbook(); ws = wb.active; ws.title = "BangKe"
F = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
thin = Side(style="thin"); dbl = Side(style="double")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
grid = Border(left=thin, right=thin, bottom=thin)
left = Alignment(horizontal="left", vertical="center")
center = Alignment(horizontal="center", vertical="center", wrap_text=True)
right = Alignment(horizontal="right", vertical="center")
for col, w in zip("ABCDEFGH", [6, 12, 24, 12, 12, 10, 14, 8]):
    ws.column_dimensions[col].width = w

def put(ref, val, font=None, align=None, border=None, fill=None, fmt=None, merge=None):
    c = ws[ref]; c.value = val; c.font = font or F()
    if align: c.alignment = align
    if border: c.border = border
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if merge: ws.merge_cells(f"{ref}:{merge}")

# logo A1 (2 cột x 4 dòng), tiêu đề bên phải
put("A1", "{{*image:logo}}")
ws.row_dimensions[2].height = 26
put("C2", "BẢNG KÊ CHUYẾN ĐI", F(size=16, bold=True), center, merge="H2")
put("C3", "Ngày in: {{printed_on}}", F(italic=True), center, merge="H3")

rows = [(6, "Khách hàng:", "{{customer}}"),
        (7, "Địa chỉ:", "{{address}}"),
        (8, "Email:", "{{email}}")]
for r, label, value in rows:
    put(f"A{r}", label, F(bold=True), merge=f"B{r}")
    put(f"C{r}", value, align=left, merge=f"H{r}")

ws.row_dimensions[10].height = 28
hdr = PatternFill("solid", fgColor="DDEBF7")
for col, title in zip("ABCDEFGH", ["STT", "Travel", "Đại lý", "Từ ngày", "Đến ngày",
                                   "Trạng thái", "Tổng tiền", "Tiền tệ"]):
    put(f"{col}10", title, F(bold=True), center, box, hdr)
# dòng mẫu: mở {{#travels}} ở ô đầu, đóng {{/travels}} ở ô cuối -> nhân dòng mỗi travel
line = [("A", "{{#travels}}{{@index}}", center, None), ("B", "{{travel_id}}", center, None),
        ("C", "{{agency}}", None, None), ("D", "{{begin_date}}", center, "dd/mm/yyyy"),
        ("E", "{{end_date}}", center, "dd/mm/yyyy"), ("F", "{{status}}", center, None),
        ("G", "{{total_price}}", None, "#,##0.00"), ("H", "{{currency}}{{/travels}}", center, None)]
for col, value, align, fmt in line:
    put(f"{col}11", value, align=align, border=grid, fmt=fmt)

put("A12", "Số chuyến đi: {{cnt:travels}}", merge="E12")
put("F12", "Tổng cộng:", F(bold=True), right)
put("G12", "{{sum:travels.total_price}}", F(size=11, bold=True), border=Border(top=dbl), fmt="#,##0.00")

ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = "landscape"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0

xlsx = out / "ZDMO01_template.xlsx"
wb.save(xlsx)
# Target tương đối như Excel (engine không đọc được "/xl/...")
buf = io.BytesIO()
with zipfile.ZipFile(xlsx) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "xl/_rels/workbook.xml.rels":
            data = data.replace(b'Target="/xl/', b'Target="')
        zout.writestr(item, data)
xlsx.write_bytes(buf.getvalue())

sample = {
    "printed_on": "2026-10-14",
    "customer": "Sample Customer",
    "address": "1 Sample Street, 10000 Sample City, DE",
    "email": "customer@example.com",
    "travels": [
        {"travel_id": "00000001", "agency": "Sunshine Travel", "begin_date": "2026-10-20",
         "end_date": "2026-10-25", "status": "B", "total_price": 2435.00, "currency": "EUR"},
        {"travel_id": "00000002", "agency": "Fly High", "begin_date": "2026-11-05",
         "end_date": "2026-11-10", "status": "N", "total_price": 1180.50, "currency": "EUR"}]}
(out / "ZDMO01_sample.json").write_text(json.dumps(sample, ensure_ascii=False, indent=2), encoding="utf-8")
print("ok")
